"""Тесты листа самопроверки: проверка решений, порядок разделов, персональные данные, структура .docx."""
import sys, os, io, json, zipfile, base64
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'web', 'py'))
from lxml import etree
import normokontrol_sheet as H

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def item(i, quote='Фрагмент текста', head='1.1 Раздел', page=5, examples=None, msg='нарушение'):
    return dict(id=i, msg=msg, quote=quote, head=head, page=page, examples=examples or [])


def payload(decisions, **kw):
    p = dict(version='0.10', date_ru='25.10.2026', date_iso='2026-10-25', offset=0, pages_known=True, include_quotes=True,
             self_checks=['Титульный лист', 'Задание на ВКР'],
             groups=[dict(rule='H5', sub=None, title='Заголовок раздела', level='fix', source='ГОСТ 7.32-2017, п. 6.2.3',
                          items=[item('0-0'), item('0-1', quote='Второй')]),
                     dict(rule='LS1', sub=None, title='Перечисления', level='check', source='ГОСТ 7.32-2017, п. 6.4.6',
                          items=[item('1-0', quote='Перечисление «1.»'), item('1-1')])],
             decisions=decisions)
    p.update(kw)
    return p


GOOD = {'d': 'keep', 'clause': 'ГОСТ 7.32-2017, п. 6.4.6', 'reason': 'Так требует кафедра для этого раздела'}
ALL_FIXED = {k: {'d': 'fixed'} for k in ('0-0', '0-1', '1-0', '1-1')}


def doc_text(data):
    z = zipfile.ZipFile(io.BytesIO(data))
    x = etree.fromstring(z.read('word/document.xml'))
    return ' '.join(t.text or '' for t in x.iter(W + 't')), x, z


def test_validate():
    assert H.validate(payload({})) == dict(undecided=['0-0', '0-1', '1-0', '1-1'], bad_keep=[])
    d = dict(ALL_FIXED); d['1-0'] = {'d': 'keep', 'clause': '', 'reason': 'Достаточно длинная причина'}
    assert H.validate(payload(d))['bad_keep'] == ['1-0']
    d['1-0'] = {'d': 'keep', 'clause': 'п. раздела', 'reason': 'Достаточно длинная причина'}      # нет цифры
    assert H.validate(payload(d))['bad_keep'] == ['1-0']
    d['1-0'] = {'d': 'keep', 'clause': 'п. 6.4.6', 'reason': 'коротко'}
    assert H.validate(payload(d))['bad_keep'] == ['1-0']
    d['1-0'] = GOOD
    assert H.validate(payload(d)) == dict(undecided=[], bad_keep=[])
    d['0-0'] = dict(GOOD)                                    # «не согласен» с замечанием уровня «Исправьте» — те же требования
    assert H.validate(payload(d))['bad_keep'] == []
    d['0-0'] = {'d': 'keep', 'clause': '', 'reason': ''}
    assert H.validate(payload(d))['bad_keep'] == ['0-0']


def test_refuses_without_reason():
    r = json.loads(H.make(json.dumps(payload({}))))
    assert r['ok'] is False and r['undecided']
    d = dict(ALL_FIXED); d['1-1'] = {'d': 'keep', 'clause': 'п. 6.4.6', 'reason': ''}
    r = json.loads(H.make(json.dumps(payload(d))))
    assert r['ok'] is False and r['bad_keep'] == ['1-1'] and 'b64' not in r


def build(p):
    r = json.loads(H.make(json.dumps(p)))
    assert r['ok'], r
    return base64.b64decode(r['b64'])


def test_kept_first_and_content():
    d = dict(ALL_FIXED); d['1-0'] = GOOD
    text, x, z = doc_text(build(payload(d)))
    assert text.index('1. Оставлено без изменений (1)') < text.index('2. Исправлено студентом (3)')
    i1, i2 = text.index('Перечисление'), text.index('Второй')
    assert i1 < i2                                   # оставленное выше исправленного, хотя группа «Перечисления» идёт после
    assert 'Так требует кафедра' in text and 'ГОСТ 7.32-2017, п. 6.4.6' in text
    assert 'Всего замечаний: 4. Оставлено без изменений: 1. Исправлено: 3.' in text
    assert 'стр. 5' in text and 'раздел «1.1 Раздел»' in text
    assert 'ФИО' in text                              # в тексте листа сказано, что ФИО нет


def test_page_offset_and_unknown_pages():
    d = dict(ALL_FIXED); d['1-0'] = GOOD
    assert 'стр. 7' in doc_text(build(payload(d, offset=2)))[0]
    t = doc_text(build(payload(d, pages_known=False)))[0]
    assert 'стр.' not in t.replace('страница', '') and 'нет меток страниц Word' in t


def test_no_personal_data_in_package():
    d = dict(ALL_FIXED); d['1-0'] = dict(GOOD, reason='Пишите на ivanov@mail.ru или +7 (911) 123-45-67, паспорт 4012 345678')
    p = payload(d)
    p['groups'][1]['items'][0]['quote'] = 'Контакт: petrov@site.ru 8 911 000 11 22'
    p['file_name'] = 'Иванов_Иван_ВКР.docx'          # лишние поля в данных не должны попасть в файл
    data = build(p)
    z = zipfile.ZipFile(io.BytesIO(data))
    blob = b''.join(z.read(n) for n in z.namelist()).decode('utf8')
    for bad in ('ivanov@mail.ru', 'petrov@site.ru', '123-45-67', '911 000 11 22', '4012 345678', 'Иванов_Иван'):
        assert bad not in blob, bad
    core = z.read('docProps/core.xml').decode('utf8')
    assert 'creator' not in core and 'lastModifiedBy' not in core
    assert '[e-mail]' in blob and '[телефон]' in blob


def test_quotes_can_be_excluded():
    d = dict(ALL_FIXED); d['1-0'] = GOOD
    p = payload(d, include_quotes=False)
    p['groups'][1]['items'][0]['msg'] = 'нумерация без скобки: «1.»'
    t = doc_text(build(p))[0]
    assert 'Перечисление' not in t and 'Второй' not in t and 'нумерация без скобки: «…»' in t


def test_invalid_xml_chars_and_structure():
    d = dict(ALL_FIXED); d['1-0'] = dict(GOOD, reason='Причина с управляющим символом \x01 и переносом\nстроки')
    p = payload(d); p['groups'][0]['items'][0]['quote'] = 'A\x0bB\ud800C'
    data = build(p)
    z = zipfile.ZipFile(io.BytesIO(data))
    assert z.testzip() is None and z.namelist()[0] == '[Content_Types].xml'
    for n in z.namelist():
        etree.fromstring(z.read(n))                  # все части — корректный XML
    t = ' '.join(x.text or '' for x in etree.fromstring(z.read('word/document.xml')).iter(W + 't'))
    assert 'и переносом строки' in t
    assert '\x01' not in t and '\x0b' not in t
    ct = z.read('[Content_Types].xml').decode()
    for part in ('/word/document.xml', '/word/styles.xml', '/word/footer1.xml', '/docProps/core.xml'):
        assert part in ct


def test_no_findings():
    p = payload({}); p['groups'] = []
    t = doc_text(build(p))[0]
    assert 'Всего замечаний: 0' in t and 'Замечаний, оставленных без изменений, нет.' in t


def test_python_docx_reads_it():
    try:
        import docx
    except ImportError:
        return
    d = dict(ALL_FIXED); d['1-0'] = GOOD
    doc = docx.Document(io.BytesIO(build(payload(d))))
    assert len(doc.tables) == 2 and len(doc.tables[0].rows) == 2 and len(doc.tables[1].rows) == 4
    assert doc.core_properties.author == ''


if __name__ == '__main__':
    n = 0
    for k, v in sorted(globals().items()):
        if k.startswith('test_') and callable(v):
            v(); n += 1; print('ok', k)
    print('тестов пройдено:', n)
